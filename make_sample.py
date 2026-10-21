from docx import Document
from docx.shared import Inches, Pt

doc = Document()
doc.add_heading('Quarterly Equipment Schedule', level=1)
doc.add_paragraph('Vendor: Northline Industrial Supply — account 4471-B')

rows = [
    ('Item code', 'Description', 'Qty', 'Unit price', 'Line total'),
    ('NX-1042', 'Hydraulic hose, 3/4 in, 5 m', '12', '£48.50', '£582.00'),
    ('NX-1043', 'Hydraulic hose, 1 in, 5 m', '8', '£61.20', '£489.60'),
    ('BR-2210', 'Bearing kit, sealed', '40', '£12.75', '£510.00'),
    ('BR-2211', 'Bearing kit, open', '25', '£9.40', '£235.00'),
    ('FL-3301', 'Filter element, 10 micron', '60', '£7.15', '£429.00'),
    ('FL-3302', 'Filter element, 25 micron', '60', '£6.80', '£408.00'),
    ('SV-1180', 'Servo valve, refurbished', '3', '£742.00', '£2,226.00'),
    ('SV-1181', 'Servo valve, new', '1', '£2,115.00', '£2,115.00'),
    ('', '', '', 'Subtotal', '£6,994.60'),
    ('', '', '', 'VAT (20%)', '£1,398.92'),
    ('', '', '', 'Total', '£8,393.52'),
]

table = doc.add_table(rows=len(rows), cols=5)
table.style = 'Table Grid'
for r, values in enumerate(rows):
    for c, value in enumerate(values):
        cell = table.cell(r, c)
        cell.text = value
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)

doc.add_page_break()
doc.add_heading('Delivery notes', level=2)
doc.add_paragraph('All items shipped from Leeds depot on 14 March 2026.')

table2 = doc.add_table(rows=4, cols=4)
table2.style = 'Table Grid'
rows2 = [
    ('Pallet', 'Contents', 'Weight (kg)', 'Signed by'),
    ('P-01', 'Hoses and fittings', '84', 'R. Okafor'),
    ('P-02', 'Bearings and filters', '112', 'R. Okafor'),
    ('P-03', 'Servo valves (crated)', '46', 'M. Ellery'),
]
for r, values in enumerate(rows2):
    for c, value in enumerate(values):
        table2.cell(r, c).text = value

doc.save('data/sample_schedule.docx')
print('docx written')
